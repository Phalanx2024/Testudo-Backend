import os
import sys
import pandas as pd
from datetime import datetime
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from database.db_controller import DatabaseController
import re


def load_sheet(filepath: str, sheet_name: str):
    df = pd.read_excel(filepath, sheet_name=sheet_name, engine='openpyxl', header=None)
    # find header row
    expected = {'date', 'company'}
    header_row = None
    for i in range(min(len(df), 100)):
        vals = set(str(v).strip().lower() for v in df.iloc[i].tolist())
        if expected.issubset(vals):
            header_row = i
            break
    if header_row is None:
        # Excel data starts at row 25 => header likely row 24 (index 23)
        header_row = 23
    headers = [str(c).strip() for c in df.iloc[header_row].tolist()]
    df = df.iloc[header_row + 1 :].copy()
    df.columns = headers
    return df, header_row


def build_colmap(df):
    lower_cols = {c.lower().strip(): c for c in df.columns}
    def fm(opts):
        for o in opts:
            if o in lower_cols:
                return lower_cols[o]
        return None
    return {
        'Date': fm(['date']),
        'Company': fm(['company']),
        'Link': fm(['link', 'links', 'url', 'report link', 'report url'])
    }


def extract_links(filepath: str, df: pd.DataFrame, header_row: int, sheet_name: str, link_col_name: str, link_col_idx_override: int = None):
    wb = load_workbook(filepath, data_only=False, read_only=False)
    ws = wb[sheet_name]
    if link_col_idx_override:
        link_col_idx = link_col_idx_override
    else:
        link_col_idx = df.columns.get_loc(link_col_name) + 1  # 1-based
    links = []
    for i in range(len(df)):
        sheet_row_num = header_row + 2 + i
        cell = ws.cell(row=sheet_row_num, column=link_col_idx)
        href = None
        if cell is not None and getattr(cell, 'hyperlink', None) is not None:
            href = getattr(cell.hyperlink, 'target', None)
        if not href and isinstance(cell.value, str):
            text = cell.value
            if text.startswith(('http://','https://')):
                href = text
            else:
                m = re.search(r'(https?://[^\s\)\]]+)', text, flags=re.IGNORECASE)
                if m:
                    href = m.group(1)
                else:
                    m2 = re.search(r'(www\.[^\s\)\]]+)', text, flags=re.IGNORECASE)
                    if m2:
                        href = 'https://' + m2.group(1)
        links.append(href)
    return links


def main():
    if len(sys.argv) < 2:
        print('Usage: python fix_muddy_waters_links.py "/path/to/SSA Report Database.xlsx" [--sheet SHEET_NAME] [--seller SELLER_NAME] [--link-col J]')
        sys.exit(1)
    excel_path = sys.argv[1]
    link_col_letter = None
    sheet_name = 'Muddy Waters'
    seller_name = 'Muddy Waters Research'
    # parse flags
    i = 2
    while i < len(sys.argv):
        if sys.argv[i] == '--link-col' and i + 1 < len(sys.argv):
            link_col_letter = sys.argv[i + 1].strip().upper()
            i += 2
        elif sys.argv[i] == '--sheet' and i + 1 < len(sys.argv):
            sheet_name = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == '--seller' and i + 1 < len(sys.argv):
            seller_name = sys.argv[i + 1]
            i += 2
        else:
            i += 1
    if not os.path.isfile(excel_path):
        print(f'File not found: {excel_path}')
        sys.exit(1)

    load_dotenv()
    df, header_row = load_sheet(excel_path, sheet_name)
    colmap = build_colmap(df)
    if not colmap['Company'] or not colmap['Date']:
        print('Missing required columns (Company/Date) in sheet')
        sys.exit(1)

    link_idx_override = None
    if link_col_letter:
        try:
            link_idx_override = column_index_from_string(link_col_letter)
            print(f"Using explicit link column: {link_col_letter} (index {link_idx_override})")
        except Exception as e:
            print(f"Invalid --link-col value '{link_col_letter}': {e}")
            sys.exit(1)

    if not colmap['Link'] and not link_idx_override:
        print('Link column not detected and no --link-col provided. Please specify --link-col (e.g., J).')
        sys.exit(1)

    links = extract_links(excel_path, df, header_row, sheet_name, colmap['Link'] or '', link_idx_override)
    df = df.reset_index(drop=True).copy()
    df['__extracted_link'] = links

    # normalize date and title key
    def norm_date(x):
        if pd.isna(x):
            return None
        if isinstance(x, (datetime, pd.Timestamp)):
            return x.date().isoformat()
        try:
            return pd.to_datetime(str(x)).date().isoformat()
        except Exception:
            return None

    df['__date'] = df[colmap['Date']].apply(norm_date)
    df['__company'] = df[colmap['Company']].astype(str).str.strip()
    df['__report_title'] = df['__company'] + ' - ' + df['__date'].fillna('')

    db = DatabaseController()
    # fetch current titles and ids
    rows = db.db.execute_query(
        "SELECT id, target_company, publication_date, link FROM short_reports WHERE short_seller = %s",
        params=(seller_name,),
        fetch=True
    )
    # Build indices
    by_company_date = {}
    by_company = {}
    for r in rows or []:
        rid, company, pub_date, curr_link = r
        key = (str(company or '').strip().lower(), pub_date.isoformat() if pub_date else None)
        by_company_date[key] = (rid, curr_link)
        comp_key = str(company or '').strip().lower()
        by_company.setdefault(comp_key, []).append((rid, pub_date, curr_link))

    # Build excel map by company
    from collections import defaultdict
    excel_by_company = defaultdict(list)
    for _, er in df.iterrows():
        comp = str(er['__company'] or '').strip().lower()
        d = er['__date']
        href = er['__extracted_link']
        if comp and href and str(href).startswith(('http://','https://')):
            excel_by_company[comp].append((d, href))

    updates = 0
    skipped = 0
    unchanged = 0
    # Iterate DB rows and choose best excel link per row
    for comp_key, candidates in by_company.items():
        excel_rows = excel_by_company.get(comp_key, [])
        if not excel_rows:
            # No excel link found for this company
            for rid, pub_date, current_link in candidates:
                print(f"Skip: no excel link for company='{comp_key}', id={rid}, date={pub_date}, current={current_link}")
            skipped += len(candidates)
            continue
        for rid, pub_date, current_link in candidates:
            # Determine date key
            db_date = pub_date.isoformat() if pub_date else None
            # Prefer exact date match
            href = None
            if db_date:
                for d, link in excel_rows:
                    if d == db_date:
                        href = link
                        break
            # If no exact match and only one excel row, use it
            if not href and len(excel_rows) == 1:
                href = excel_rows[0][1]

            # If still no href, skip
            if not href:
                print(f"Skip: no matching href for company='{comp_key}', id={rid}, date={pub_date}")
                skipped += 1
                continue

            # Update only if DB link missing/non-http or different
            if current_link and str(current_link).startswith(('http://','https://')) and current_link == href:
                print(f"Unchanged: company='{comp_key}', id={rid}, date={pub_date}, link={current_link}")
                unchanged += 1
                continue

            print(f"Update: company='{comp_key}', id={rid}, date={pub_date}, current={current_link} -> new={href}")
            ok = db.db.execute_query("UPDATE short_reports SET link = %s WHERE id = %s", params=(href, rid))
            if ok:
                updates += 1

    print(f'Links updated: {updates}, unchanged: {unchanged}, skipped: {skipped}')


if __name__ == '__main__':
    main()


