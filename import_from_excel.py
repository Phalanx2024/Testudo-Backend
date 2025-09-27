import os
import sys
import pandas as pd
from datetime import datetime
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
import re

from database.db_controller import DatabaseController


def load_sheet(filepath: str, sheet_name: str):
    df = pd.read_excel(filepath, sheet_name=sheet_name, engine='openpyxl', header=None)
    expected = {'date', 'company'}
    header_row = None
    for i in range(min(len(df), 100)):
        vals = set(str(v).strip().lower() for v in df.iloc[i].tolist())
        if expected.issubset(vals):
            header_row = i
            break
    if header_row is None:
        header_row = 23
    headers = [str(c).strip() for c in df.iloc[header_row].tolist()]
    df = df.iloc[header_row + 1 :].copy()
    df.columns = headers
    return df, header_row


def build_colmap(df):
    lower_cols = {c.lower().strip(): c for c in df.columns}
    def fm(opts):
        for o in opts:
            if o in lower_cols:
                return lower_cols[o]
        return None
    return {
        'Date': fm(['date']),
        'Company': fm(['company']),
        'Ticker': fm(['ticker']),
        'Sector': fm(['sector']),
        'Link': fm(['link', 'links', 'url', 'report link', 'report url'])
    }


def extract_links(filepath: str, sheet_name: str, header_row: int, df: pd.DataFrame, link_col_name: str, link_col_letter: str = None):
    wb = load_workbook(filepath, data_only=False, read_only=False)
    ws = wb[sheet_name]
    if link_col_letter:
        link_col_idx = column_index_from_string(link_col_letter)
    else:
        link_col_idx = df.columns.get_loc(link_col_name) + 1
    links = []
    for i in range(len(df)):
        sheet_row_num = header_row + 2 + i
        cell = ws.cell(row=sheet_row_num, column=link_col_idx)
        href = None
        if cell is not None and getattr(cell, 'hyperlink', None) is not None:
            href = getattr(cell.hyperlink, 'target', None)
        if not href and isinstance(cell.value, str):
            text = cell.value
            if text.startswith(('http://','https://')):
                href = text
            else:
                m = re.search(r'(https?://[^\s\)\]]+)', text, flags=re.IGNORECASE)
                if m:
                    href = m.group(1)
                else:
                    m2 = re.search(r'(www\.[^\s\)\]]+)', text, flags=re.IGNORECASE)
                    if m2:
                        href = 'https://' + m2.group(1)
        links.append(href)
    return links


def norm_date(x):
    if pd.isna(x):
        return None
    if isinstance(x, (datetime, pd.Timestamp)):
        return x.date()
    try:
        return pd.to_datetime(str(x)).date()
    except Exception:
        return None


def main():
    if len(sys.argv) < 5:
        print('Usage: python import_from_excel.py "/path/to/SSA Report Database.xlsx" --sheet SHEET --seller SELLER [--link-col J]')
        sys.exit(1)
    excel_path = sys.argv[1]
    sheet_name = None
    seller_name = None
    link_col_letter = None

    i = 2
    while i < len(sys.argv):
        if sys.argv[i] == '--sheet' and i + 1 < len(sys.argv):
            sheet_name = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == '--seller' and i + 1 < len(sys.argv):
            seller_name = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == '--link-col' and i + 1 < len(sys.argv):
            link_col_letter = sys.argv[i + 1].strip().upper()
            i += 2
        else:
            i += 1

    if not os.path.isfile(excel_path):
        print(f'File not found: {excel_path}')
        sys.exit(1)
    if not sheet_name or not seller_name:
        print('Both --sheet and --seller are required')
        sys.exit(1)

    load_dotenv()
    df, header_row = load_sheet(excel_path, sheet_name)
    colmap = build_colmap(df)
    if not colmap['Company'] or not colmap['Date']:
        print('Missing required columns (Company/Date) in sheet')
        sys.exit(1)

    # Extract links
    if not colmap['Link'] and not link_col_letter:
        print('Link column not detected; please provide --link-col')
        sys.exit(1)
    links = extract_links(excel_path, sheet_name, header_row, df, colmap['Link'] or '', link_col_letter)

    df = df.reset_index(drop=True).copy()
    df['__link'] = links
    df['__date'] = df[colmap['Date']].apply(norm_date)

    # Build records
    records = []
    for _, row in df.iterrows():
        company = str(row[colmap['Company']]).strip() if pd.notna(row.get(colmap['Company'])) else ''
        ticker = str(row[colmap['Ticker']]).strip() if colmap['Ticker'] and pd.notna(row.get(colmap['Ticker'])) else ''
        sector = str(row[colmap['Sector']]).strip() if colmap['Sector'] and pd.notna(row.get(colmap['Sector'])) else ''
        link = row['__link'] if pd.notna(row['__link']) else ''
        pub_date = row['__date']

        if not link:
            continue

        report_title = f"{company} - {pub_date.isoformat()}" if company and pub_date else company or link

        rec = {
            'publication_date': pub_date,
            'report_title': report_title,
            'link': link,
            'target_company': company,
            'short_seller': seller_name,
            'ticker': ticker,
            'sector': sector,
            'last_update': datetime.utcnow().date(),
        }
        records.append(rec)

    # Dedupe in-batch by link
    seen = set()
    unique_records = []
    for rec in records:
        if rec['link'] in seen:
            continue
        seen.add(rec['link'])
        unique_records.append(rec)

    db = DatabaseController()
    # Check existing by link for this seller
    existing = db.db.execute_query("SELECT link FROM short_reports WHERE short_seller = %s", params=(seller_name,), fetch=True)
    existing_links = set(r[0] for r in (existing or []) if r and r[0])
    to_insert = [r for r in unique_records if r['link'] not in existing_links]

    inserted = 0
    for rec in to_insert:
        if db.insert_data('short_reports', rec):
            inserted += 1

    print(f"Prepared {len(records)} records, {len(unique_records)} unique by link, {len(to_insert)} new. Inserted: {inserted}.")


if __name__ == '__main__':
    main()


